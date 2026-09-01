# -*- coding: utf-8 -*-
"""v5 -> v6 · rehace la diapo 25 sin el contraste Cap. 6 / Cap. 8:
   pregunta -> características del modelo -> diseño -> consideraciones."""
from pptx import Presentation
from pptx.util import Pt, Cm
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

SRC='/sessions/dreamy-focused-thompson/mnt/PAC_v2/tesis_cap/Tesis_PAC_presentacion_conceptual_Mac_editable_v5.pptx'
DST='/tmp/Tesis_PAC_presentacion_conceptual_Mac_editable_v6.pptx'
NB=chr(0x00A0)
BLUE=RGBColor(0x00,0x70,0xC0); INK=RGBColor(0x1E,0x25,0x2D); GREY=RGBColor(0x7A,0x87,0x94)
CARDBG=RGBColor(0xF5,0xF8,0xFA); CARDLN=RGBColor(0xDF,0xE7,0xED); F='Aptos'

p=Presentation(SRC); s=p.slides[24]

# fuera todo menos plantilla (fondo, barra, pie, titulo)
for sh in list(s.shapes):
    if sh.name not in ('Rectangle 1','Rectangle 2','TextBox 3','TextBox 4'):
        sh._element.getparent().remove(sh._element)

t=[sh for sh in s.shapes if sh.name=='TextBox 4'][0]
runs=list(t.text_frame.paragraphs[0].runs)
runs[0].text='El modelo: qué pregunta responde y con qué información'
for r in runs[1:]: r.text=''

# linea de pregunta, con el estilo QuestionLine del mazo
q=s.shapes.add_textbox(Cm(1.40),Cm(2.30),Cm(30.99),Cm(1.40))
q.text_frame.word_wrap=True
pr=q.text_frame.paragraphs[0]; pr.alignment=PP_ALIGN.CENTER
r=pr.add_run(); r.text='Pregunta: '
r.font.size=Pt(23); r.font.bold=True; r.font.italic=True; r.font.name=F; r.font.color.rgb=GREY
r=pr.add_run(); r.text=f'¿habrá un evento severo en los próximos H{NB}minutos?'
r.font.size=Pt(23); r.font.italic=True; r.font.name=F; r.font.color.rgb=GREY

def card(l,t,w,h,titulo,bullets,sz=12.5):
    c=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Cm(l),Cm(t),Cm(w),Cm(h))
    c.fill.solid(); c.fill.fore_color.rgb=CARDBG
    c.line.color.rgb=CARDLN; c.line.width=Pt(1); c.shadow.inherit=False
    tx=s.shapes.add_textbox(Cm(l+0.55),Cm(t+0.32),Cm(w-1.10),Cm(h-0.7))
    tf=tx.text_frame; tf.word_wrap=True
    pr=tf.paragraphs[0]; r=pr.add_run(); r.text=titulo
    r.font.size=Pt(15); r.font.bold=True; r.font.name=F; r.font.color.rgb=BLUE
    for b in bullets:
        pr=tf.add_paragraph(); pr.space_before=Pt(8)
        r=pr.add_run(); r.text='• '+b
        r.font.size=Pt(sz); r.font.name=F; r.font.color.rgb=INK

T,H,W,G=4.20,8.90,9.96,0.60
card(1.40,T,W,H,'Características del modelo',
     ['LightGBM, 300 árboles; regresión logística como referencia.',
      'Cinco bloques de variables: estado PAC multiescala, posición en la noche, historial reciente de eventos, morfología reciente e historial acumulado.',
      'Todas miran hacia atrás: ninguna usa información posterior al cierre de la ventana.'])
card(1.40+W+G,T,W,H,'Diseño del análisis',
     ['Cohorte estricta: 8 pacientes, 540 noches.',
      f'Ventanas de 30{NB}s contiguas, sin solapamiento: 452.955 en total.',
      f'Horizontes H ∈ {{2, 5, 10, 15}}{NB}min, fijados de antemano.',
      'LOPO-CV: el paciente entero queda fuera del entrenamiento.'])
card(1.40+2*(W+G),T,W,H,'Consideraciones',
     ['La salida es un ranking de riesgo, no una probabilidad calibrada.',
      'La prevalencia del evento varía 20× entre pacientes.',
      'Los Estados PAC se entrenaron sobre esta misma cohorte.',
      'Sin validación externa; n = 8.'])
p.save(DST)
q2=Presentation(DST)
print('slides:',len(q2.slides))
for sh in q2.slides[24].shapes:
    if sh.has_text_frame and sh.text_frame.text.strip():
        print('  ·',sh.text_frame.text.replace(chr(10),' | ')[:95])
