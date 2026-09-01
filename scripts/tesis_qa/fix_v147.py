# -*- coding: utf-8 -*-
"""v146 -> v147: cierre APA 7 (mayuscula tras dos puntos, descriptores de software),
normalizacion de '(Anexo Suplementario)' en los indices y reubicacion del titulo de la Tabla 1.3.

Uso: python3 scripts/tesis_qa/fix_v147.py
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.text.run import Run

sys.path.insert(0, str(Path(__file__).parent))
from lib_docx import para_replace  # noqa: E402


class ParrafoProfundo:
    """Vista de un parrafo cuyos runs incluyen los que estan dentro de hyperlinks.

    `p.runs` de python-docx solo devuelve los w:r hijos directos; en los indices
    de figuras y tablas el texto vive dentro de un w:hyperlink y queda invisible.
    """

    def __init__(self, p):
        self._p = p
        self.runs = [Run(r, p) for r in p._p.iter(qn("w:r"))]

SRC = Path("tesis_cap/Tesis_PAC_v146.docx")
DST = Path("tesis_cap/Tesis_PAC_v147.docx")

RE_ENTRADA_APA = re.compile(r"^[A-ZÁÉÍÓÚÑ][^()]{0,300}\(\d{4}[a-z]?\)\.")
RE_TRAS_DOSPUNTOS = re.compile(r"(\(\d{4}[a-z]?\)\.\s+[^.]*?:\s+)([a-z])")

# Descriptor APA entre corchetes para entradas de software / modelos.
DESCRIPTORES = {
    "https://streamlit.io": "[Software]",
    "https://plot.ly": "[Software]",
    "https://www.anthropic.com": "[Modelo de lenguaje grande]",
}


def bloque_referencias(doc):
    """Indice del parrafo 'Referencias'; las entradas van despues."""
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip() == "Referencias":
            return i
    raise RuntimeError("no se encontro el bloque de Referencias")


def fix_mayuscula_tras_dospuntos(doc, inicio):
    """APA 7: la primera palabra despues de dos puntos en un titulo va en mayuscula."""
    n = 0
    for p in doc.paragraphs[inicio + 1:]:
        texto = p.text.strip()
        if not RE_ENTRADA_APA.match(texto):
            continue
        m = RE_TRAS_DOSPUNTOS.search(texto)
        if not m:
            continue
        ini = max(0, m.start(1) - 12)
        viejo = texto[ini:m.end(2)]
        nuevo = viejo[:-1] + m.group(2).upper()
        if para_replace(p, viejo, nuevo):
            n += 1
            print(f"  APA mayuscula: ...{nuevo[-60:]}")
    return n


def fix_descriptores_software(doc, inicio):
    """APA 7: software y modelos llevan descriptor entre corchetes tras el titulo."""
    n = 0
    for p in doc.paragraphs[inicio + 1:]:
        texto = p.text.strip()
        for url, desc in DESCRIPTORES.items():
            if not texto.endswith(url) or desc in texto:
                continue
            if para_replace(p, f". {url}", f" {desc}. {url}"):
                n += 1
                print(f"  descriptor {desc} -> {texto[:40]}")
    return n


def fix_anexo_suplementario(doc):
    """Una sola forma en los indices: '<titulo>. (Anexo Suplementario)'."""
    n = 0
    for p in doc.paragraphs:
        texto = p.text
        if "(Anexo Suplementario)" not in texto or "\t" not in texto:
            continue
        m = re.search(r"([^\t]*?) ?\(Anexo Suplementario\) ?(\.?)", texto)
        if not m:
            continue
        viejo = m.group(0)
        nuevo = m.group(1).rstrip(" .") + ". (Anexo Suplementario)"
        if viejo != nuevo and para_replace(ParrafoProfundo(p), viejo, nuevo):
            n += 1
            print(f"  indice: {nuevo[:70]}")
    return n


def fix_titulo_tabla_1_3(doc):
    """APA: el titulo de la tabla va inmediatamente encima; la prosa, antes del titulo."""
    ps = doc.paragraphs
    for i, p in enumerate(ps):
        if not p.text.strip().startswith("Tabla 1.3. Tabla de trazabilidad"):
            continue
        if "\t" in p.text:  # entrada del indice, no el titulo real
            continue
        siguiente = ps[i + 1]
        if not siguiente.text.strip().startswith("La siguiente tabla traza"):
            return 0
        p._p.addprevious(siguiente._p)  # la prosa pasa por encima del titulo
        print("  Tabla 1.3: prosa movida por encima del titulo")
        return 1
    return 0


def main():
    doc = Document(str(SRC))
    inicio = bloque_referencias(doc)

    print("1. APA 7 - mayuscula tras dos puntos")
    a = fix_mayuscula_tras_dospuntos(doc, inicio)
    print("2. APA 7 - descriptores de software")
    b = fix_descriptores_software(doc, inicio)
    print("3. Indices - '(Anexo Suplementario)'")
    c = fix_anexo_suplementario(doc)
    print("4. Tabla 1.3 - titulo pegado a la tabla")
    d = fix_titulo_tabla_1_3(doc)

    doc.save(str(DST))
    print(f"\n{a} titulos APA · {b} descriptores · {c} entradas de indice · {d} reubicacion")
    print(f"guardado: {DST}")


if __name__ == "__main__":
    main()
