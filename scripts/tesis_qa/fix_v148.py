# -*- coding: utf-8 -*-
"""v147 -> v148: sincroniza los indices de figuras y tablas con los pies del cuerpo.

Los indices de figuras y tablas NO son campos TOC (ver CLAUDE.md §27): son parrafos
escritos a mano con un PAGEREF al final. F9 les actualiza solo el numero de pagina.
Por eso quedaron ahi residuos de reemplazos globales que la auditoria no mira.

Uso: python3 scripts/tesis_qa/fix_v148.py
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.text.run import Run

sys.path.insert(0, str(Path(__file__).parent))
from lib_docx import para_replace  # noqa: E402

SRC = Path("tesis_cap/Tesis_PAC_v147.docx")
DST = Path("tesis_cap/Tesis_PAC_v148.docx")

RE_PIE = re.compile(r"^((?:Figura|Tabla) \d+\.\d+)\.\s*(.*)$")

# Correcciones sobre entradas del indice. Clave = numero de figura/tabla.
# Cada par es (texto viejo, texto nuevo). El resto de las divergencias
# indice/cuerpo son abreviaciones deliberadas y no se tocan.
INDICE = {
    # el titulo lo cambio el autor en el cuerpo; el indice quedo con el viejo
    "Tabla 10.2": ("Publicaciones asociadas al Proyecto PAC.",
                   "Trabajos previstos para publicación derivados del Proyecto PAC."),
    # §26: nunca 'full' / 'strict'
    "Tabla 4.7": ("cohortes full (12 p.) y strict (8 p.)",
                  "cohortes completa (12 p.) y estricta (8 p.)"),
    # §26: 'dimensiones' es concepto; las columnas de la matriz son 'variables'
    "Figura 7.10": ("perfil de dimensiones nocturnas",
                    "perfil de variables nocturnas"),
    "Figura 8.2": ("Importancia de dimensiones LightGBM",
                   "Importancia de variables LightGBM"),
    # espacios duros: el indice usa espacios normales, el cuerpo no-rompibles
    "Tabla 8.3": ("(H = 5 min) por paciente",
                  "(H = 5 min) por paciente"),
    # falta el punto final
    "Figura 6.13": ("clasificación nocturna binaria",
                    "clasificación nocturna binaria."),
}

# Correcciones sobre pies del cuerpo.
CUERPO = {
    "Tabla 4.1": ("según ODI3 (N = 560 noches)", "según ODI3 (N = 560 noches)."),
}


class ParrafoProfundo:
    """Runs incluidos los que viven dentro de un w:hyperlink (invisibles a p.runs)."""

    def __init__(self, p):
        self._p = p
        self.runs = [Run(r, p) for r in p._p.iter(qn("w:r"))]


def mapa_pies(doc):
    """{numero: (parrafo_indice, parrafo_cuerpo)} para figuras y tablas."""
    idx, cue = {}, {}
    for p in doc.paragraphs:
        texto = p.text.strip()
        m = RE_PIE.match(texto.split("\t")[0])
        if not m:
            continue
        destino = idx if "\t" in texto else cue
        destino.setdefault(m.group(1), p)
    return idx, cue


def aplicar(mapa, correcciones, profundo, etiqueta):
    n = 0
    for num, (viejo, nuevo) in correcciones.items():
        p = mapa.get(num)
        if p is None:
            print(f"  !! {num}: no encontrado en {etiqueta}")
            continue
        objetivo = ParrafoProfundo(p) if profundo else p
        if nuevo in p.text and viejo not in p.text:
            print(f"  == {num}: ya corregido")
            continue
        if para_replace(objetivo, viejo, nuevo):
            n += 1
            print(f"  {etiqueta} {num}: {nuevo[:70]}")
        else:
            print(f"  !! {num}: no se pudo reemplazar {viejo[:40]!r}")
    return n


def informe_divergencias(doc):
    """Lista lo que sigue difiriendo entre indice y cuerpo, para revision manual."""
    idx, cue = mapa_pies(doc)
    restantes = []
    for num, p in idx.items():
        if num not in cue:
            continue
        ti = re.sub(r"\s*\(Anexo Suplementario\)\s*\.?$", "",
                    RE_PIE.match(p.text.strip().split("\t")[0]).group(2)).strip()
        tc = RE_PIE.match(cue[num].text.strip()).group(2).strip()
        if ti != tc:
            restantes.append((num, ti, tc))
    return restantes


def restaurar_bordes_none(doc):
    """Word poda los bordes con val='none' al guardar, por redundantes.

    No cambia como se ve (sin tblStyle, un borde ausente no dibuja nada), pero deja
    el .docx dependiendo del estilo por defecto del renderizador. Se reponen para que
    las tres lineas APA queden declaradas de forma explicita.
    """
    n = 0
    for t in doc.tables:
        tblPr = t._tbl.find(qn("w:tblPr"))
        b = tblPr.find(qn("w:tblBorders")) if tblPr is not None else None
        if b is None:
            continue
        for tag in ("left", "right", "insideH", "insideV"):
            if b.find(qn("w:" + tag)) is not None:
                continue
            el = b.makeelement(qn("w:" + tag), {qn("w:val"): "none"})
            b.append(el)
            n += 1
    return n


def main():
    doc = Document(str(SRC))
    idx, cue = mapa_pies(doc)
    print(f"indice: {len(idx)} entradas · cuerpo: {len(cue)} pies\n")

    print("1. Entradas del indice")
    a = aplicar(idx, INDICE, profundo=True, etiqueta="idx")
    print("2. Pies del cuerpo")
    b = aplicar(cue, CUERPO, profundo=False, etiqueta="cue")
    print("3. Bordes explicitos 'none' (Word los poda al guardar)")
    print(f"  {restaurar_bordes_none(doc)} bordes repuestos en {len(doc.tables)} tablas")

    doc.save(str(DST))
    print(f"\n{a} entradas de indice · {b} pies de cuerpo")

    restantes = informe_divergencias(Document(str(DST)))
    print(f"\nDivergencias indice/cuerpo restantes: {len(restantes)}")
    print("(abreviaciones deliberadas del indice: el pie del cuerpo agrega N, cohorte o panel)")
    for num, ti, tc in sorted(restantes):
        print(f"  {num}\n    idx    {ti[:110]}\n    cuerpo {tc[:110]}")
    print(f"\nguardado: {DST}")


if __name__ == "__main__":
    main()
