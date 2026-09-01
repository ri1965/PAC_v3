# -*- coding: utf-8 -*-
"""v161 -> v162: punto 7 (terminología) y el resto del punto 2 (APA) del director.

PUNTO 7 · terminología
  · Tres usos de "RS" suelto en §6.9. La abreviatura NO se introduce en ningún lado
    del documento, así que además de violar el criterio del §26 ("siempre Risk Score
    PAC, nunca suelto") queda opaca para quien lee salteado. La cuarta aparición, en
    el glosario, define la fórmula y se deja.
  · "dentro de un Estado PAC patológico" -> minúscula: es una instancia concreta, no
    el sistema (§26). Las otras 10 apariciones en singular refieren al sistema
    (`etiquetas de Estado PAC`, `asignación de Estado PAC`) y quedan como están.

PUNTO 2 · APA en referencias
  · Terrill (2020) era la única entrada de artículo citada con URL de PubMed en vez
    de DOI, entre 51 que usan https://doi.org/. Además le faltaban volumen, número
    y páginas. Se completa: Respirology, 25(5), 475–485.
    ⚠ APA pone en cursiva el nombre de la revista Y el volumen, pero NO el número
    entre paréntesis ni las páginas: *Respirology, 25*(5), 475–485.

Uso: python3 scripts/tesis_qa/fix_v162.py <origen> <destino>
"""
import copy
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

sys.path.insert(0, str(Path(__file__).parent))
from lib_docx import iter_block_paragraphs, para_replace  # noqa: E402

SRC = Path(sys.argv[1] if len(sys.argv) > 1 else "tesis_cap/Tesis_PAC_v161.docx")
DST = Path(sys.argv[2] if len(sys.argv) > 2 else "tesis_cap/Tesis_PAC_v162.docx")

TERMINOLOGIA = [
    ("un RS alto con ODI3", "un Risk Score PAC alto con ODI3"),
    ("(RS bajo con ODI3 alto)", "(Risk Score PAC bajo con ODI3 alto)"),
    ("el RS describe el terreno", "el Risk Score PAC describe el terreno"),
    ("dentro de un Estado PAC patológico", "dentro de un estado PAC patológico"),
]

DOI_TERRILL = "https://doi.org/10.1111/resp.13635"


def arreglar_terrill(doc):
    """Completa la referencia y cambia la URL de PubMed por el DOI.

    La entrada tiene tres runs: autor+título (redonda) | 'Respirology' (cursiva) |
    resto (redonda). El volumen va en cursiva, así que se amplía el run 2; el número
    y las páginas van en redonda, así que arrancan el run 3.
    """
    for p in doc.paragraphs:
        if not p.text.strip().startswith("Terrill"):
            continue
        runs = [r for r in p.runs if r.text.strip()]
        if len(runs) < 3:
            return 0
        cursiva = next((r for r in runs if r.italic and "Respirology" in r.text), None)
        if cursiva is None or "," in cursiva.text:
            return 0                              # ya corregido
        cursiva.text = "Respirology, 25"
        ultimo = runs[-1]
        ultimo.text = f"(5), 475–485. {DOI_TERRILL}"
        return 1
    return 0


def reescalar_tablas(doc, ancho):
    """Lleva toda tabla más ancha que la caja de texto al ancho disponible.

    ⚠ Hay que tocar TRES cosas, no dos. `fix_v161` fijó `tblGrid` y el `tcW` de cada
    celda, pero dejó el `tblW` (ancho preferido de la tabla) en 9000: al abrir y
    guardar en Word, Word le hace caso a `tblW` y reexpande la grilla. El síntoma es
    que la auditoría vuelve a marcar tablas más anchas que el texto después de un
    guardado que no tocó ninguna tabla.
    """
    n = 0
    for t in doc.tables:
        grid = t._tbl.find(qn("w:tblGrid"))
        if grid is None:
            continue
        cols = [int(c.get(qn("w:w"))) for c in grid]
        if sum(cols) <= ancho:
            continue
        nuevos = [max(1, round(c * ancho / sum(cols))) for c in cols]
        nuevos[-1] += ancho - sum(nuevos)
        for col, w in zip(grid, nuevos):
            col.set(qn("w:w"), str(w))
        for fila in t.rows:
            for celda, w in zip(fila.cells, nuevos):
                tcPr = celda._tc.get_or_add_tcPr()
                tcW = tcPr.find(qn("w:tcW"))
                if tcW is None:
                    tcPr.append(tcPr.makeelement(
                        qn("w:tcW"), {qn("w:w"): str(w), qn("w:type"): "dxa"}))
                else:
                    tcW.set(qn("w:w"), str(w))
                    tcW.set(qn("w:type"), "dxa")
        tblW = t._tbl.tblPr.find(qn("w:tblW"))          # <- lo que faltaba
        if tblW is not None and tblW.get(qn("w:type")) == "dxa":
            tblW.set(qn("w:w"), str(ancho))
        n += 1
    return n


def main():
    doc = Document(str(SRC))

    print("PUNTO 7 · terminología")
    pendientes = dict(TERMINOLOGIA)
    for p in iter_block_paragraphs(doc):
        for viejo in list(pendientes):
            if viejo in p.text and para_replace(p, viejo, pendientes[viejo]):
                print(f"   {viejo!r}\n      -> {pendientes.pop(viejo)!r}")
    for viejo in pendientes:
        print(f"   !! no encontrado: {viejo!r}")

    print("\nPUNTO 2 · referencia de Terrill (2020)")
    if arreglar_terrill(doc):
        print(f"   completada con volumen, número y páginas; PubMed -> {DOI_TERRILL}")
    else:
        print("   sin cambios (¿ya corregida?)")

    s = doc.sections[0]
    ancho = s.page_width.twips - s.left_margin.twips - s.right_margin.twips
    print(f"\nPUNTO 5 · tablas reexpandidas por Word (tblW sin corregir)")
    print(f"   {reescalar_tablas(doc, ancho)} tablas devueltas al ancho de texto ({ancho} tw)")

    doc.save(str(DST))
    print(f"\nguardado: {DST}")


if __name__ == "__main__":
    main()
